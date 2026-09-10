"""Backend Flask du Chatbot Vocal.

Routes :
  GET  /health                          -> état du service
  POST /api/transcribe                  -> transcrit un fichier WAV en document Word
  GET  /api/transcriptions              -> liste l'historique des transcriptions
  GET  /api/transcriptions/<id>         -> détail d'une transcription
  GET  /api/transcriptions/<id>/download-> télécharge le document Word
  DELETE /api/transcriptions/<id>       -> supprime une transcription
  DELETE /api/transcriptions            -> vide l'historique

Configuration par variables d'environnement (voir .env.example) :
  FLASK_DEBUG, HOST, PORT, MODEL_PATH, DATA_DIR, MAX_CONTENT_LENGTH_MB, CORS_ORIGINS
"""

import json
import logging
import os
import threading
import uuid
from datetime import datetime
from tempfile import NamedTemporaryFile

from flask import Flask, jsonify, request, send_from_directory
from flask_cors import CORS
from docx import Document
import pydub  # conversion audio (nécessite ffmpeg sur le système)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
MODEL_PATH = os.environ.get("MODEL_PATH", "models/vosk-model-small-fr-0.22")
DATA_DIR = os.environ.get("DATA_DIR", "data")
TRANSCRIPTIONS_DIR = os.path.join(DATA_DIR, "transcriptions")
HISTORY_FILE = os.path.join(DATA_DIR, "transcription_history.json")
MAX_CONTENT_LENGTH_MB = int(os.environ.get("MAX_CONTENT_LENGTH_MB", "50"))
CORS_ORIGINS = [
    origin.strip()
    for origin in os.environ.get(
        "CORS_ORIGINS", "http://localhost:3000,http://127.0.0.1:3000"
    ).split(",")
    if origin.strip()
]

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger("chatbot-vocal")

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_CONTENT_LENGTH_MB * 1024 * 1024
CORS(app, resources={r"/api/*": {"origins": CORS_ORIGINS}})

os.makedirs(TRANSCRIPTIONS_DIR, exist_ok=True)
logger.info("Dossier de stockage prêt : %s", TRANSCRIPTIONS_DIR)

# Le modèle Vosk est chargé une seule fois, au premier appel (lazy loading) :
# l'API démarre même si le modèle n'est pas encore téléchargé.
_model = None
_model_lock = threading.Lock()
_history_lock = threading.Lock()


class ModelUnavailableError(RuntimeError):
    """Levée quand le modèle Vosk est absent ou impossible à charger."""


def get_model():
    """Charge (une seule fois) et retourne le modèle Vosk."""
    global _model
    if _model is None:
        with _model_lock:
            if _model is None:
                if not os.path.isdir(MODEL_PATH):
                    raise ModelUnavailableError(
                        f"Le modèle Vosk est introuvable dans '{MODEL_PATH}'. "
                        "Lancez d'abord : python scripts/download_model.py"
                    )
                try:
                    from vosk import Model

                    logger.info("Chargement du modèle Vosk depuis %s ...", MODEL_PATH)
                    _model = Model(MODEL_PATH)
                    logger.info("Modèle Vosk chargé")
                except ModelUnavailableError:
                    raise
                except Exception as exc:  # dépendance manquante, modèle corrompu...
                    raise ModelUnavailableError(
                        f"Impossible de charger le modèle Vosk : {exc}"
                    ) from exc
    return _model


# ---------------------------------------------------------------------------
# Persistance de l'historique (fichier JSON, avec verrou en attendant une BDD)
# ---------------------------------------------------------------------------
def load_history():
    if not os.path.exists(HISTORY_FILE):
        return []
    try:
        with open(HISTORY_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
            return data if isinstance(data, list) else []
    except (json.JSONDecodeError, OSError) as exc:
        logger.error("Historique illisible, repart d'une base vide : %s", exc)
        return []


def save_history(history):
    """Écriture atomique : on écrit dans un fichier temporaire puis on renomme."""
    tmp_path = f"{HISTORY_FILE}.tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(history, f, ensure_ascii=False, indent=2)
    os.replace(tmp_path, HISTORY_FILE)


# ---------------------------------------------------------------------------
# Traitement audio
# ---------------------------------------------------------------------------
def prepare_wav_for_vosk(file_path):
    """Convertit le WAV en PCM compatible Vosk : mono, 16 kHz, 16-bit."""
    audio = pydub.AudioSegment.from_file(file_path)
    audio = audio.set_channels(1).set_frame_rate(16000).set_sample_width(2)
    converted_path = file_path.replace(".wav", "_converted.wav")
    audio.export(converted_path, format="wav")
    return converted_path


def transcribe_audio(audio_path):
    """Reconnaissance vocale Vosk, retourne le texte transcrit."""
    import json as _json
    from vosk import KaldiRecognizer

    recognizer = KaldiRecognizer(get_model(), 16000)
    recognizer.SetWords(True)  # active les mots horodatés (utile plus tard)
    with open(audio_path, "rb") as wf:
        while True:
            data = wf.read(4000)
            if len(data) == 0:
                break
            recognizer.AcceptWaveform(data)
    result = _json.loads(recognizer.FinalResult())
    return result.get("text", "").strip()


def create_word_document(text, destination_path):
    doc = Document()
    doc.add_heading("Transcription Vocale", level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Date : {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}").italic = True
    doc.add_paragraph(text)
    doc.save(destination_path)


def process_audio(audio_file):
    """Sauvegarde l'audio, le transcrit, génère un document Word unique."""
    temp_audio_path = None
    converted_path = None
    try:
        # Vérifie le modèle avant tout traitement pour un message clair (503)
        get_model()
        suffix = os.path.splitext(audio_file.filename)[1] or ".wav"
        with NamedTemporaryFile(delete=False, suffix=suffix) as temp_audio:
            audio_file.save(temp_audio.name)
            temp_audio_path = temp_audio.name

        converted_path = prepare_wav_for_vosk(temp_audio_path)

        text = transcribe_audio(converted_path).strip()
        if not text:
            raise ValueError("Aucun texte transcrit (audio vide, silencieux ou inaudible).")

        # Nom de fichier unique : plus aucun écrasement entre transcriptions
        transcription_id = str(uuid.uuid4())
        timestamp = datetime.now()
        filename = (
            f"transcription_{timestamp.strftime('%Y%m%d_%H%M%S')}_"
            f"{transcription_id[:8]}.docx"
        )
        document_path = os.path.join(TRANSCRIPTIONS_DIR, filename)
        create_word_document(text, document_path)

        entry = {
            "id": transcription_id,
            "text": text,
            "filename": filename,
            "created_at": timestamp.isoformat(),
        }
        with _history_lock:
            history = load_history()
            history.append(entry)
            save_history(history)

        logger.info("Transcription %s générée (%d caractères)", transcription_id, len(text))
        return entry
    finally:
        for path in (temp_audio_path, converted_path):
            if path and os.path.exists(path):
                try:
                    os.unlink(path)
                except OSError:
                    logger.warning("Impossible de supprimer le fichier temporaire %s", path)


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/health")
def health():
    return jsonify(
        {
            "status": "ok",
            "model_path": MODEL_PATH,
            "model_available": os.path.isdir(MODEL_PATH),
        }
    )


@app.route("/api/transcribe", methods=["POST"])
def transcribe():
    if "audio" not in request.files:
        return jsonify({"error": "Aucun fichier audio reçu (champ 'audio' attendu)."}), 400

    audio_file = request.files["audio"]
    if not audio_file.filename:
        return jsonify({"error": "Aucun fichier audio sélectionné."}), 400

    if not audio_file.filename.lower().endswith(".wav"):
        return (
            jsonify({"error": "Format non supporté : seul le format WAV est accepté."}),
            415,
        )

    try:
        entry = process_audio(audio_file)
    except ModelUnavailableError as exc:
        logger.error("Modèle indisponible : %s", exc)
        return jsonify({"error": str(exc)}), 503
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 422
    except Exception:  # erreur inattendue : on logge la pile, message générique au client
        logger.exception("Erreur pendant le traitement de l'audio")
        return jsonify({"error": "Erreur interne lors du traitement de l'audio."}), 500

    entry["download_url"] = f"/api/transcriptions/{entry['id']}/download"
    return jsonify(entry), 201


@app.route("/api/transcriptions", methods=["GET"])
def list_transcriptions():
    with _history_lock:
        history = load_history()
    # Plus récentes en premier
    history = sorted(history, key=lambda item: item.get("created_at", ""), reverse=True)
    for entry in history:
        entry["download_url"] = f"/api/transcriptions/{entry['id']}/download"
    return jsonify(history)


@app.route("/api/transcriptions/<transcription_id>", methods=["GET"])
def get_transcription(transcription_id):
    with _history_lock:
        history = load_history()
    for entry in history:
        if entry["id"] == transcription_id:
            entry["download_url"] = f"/api/transcriptions/{entry['id']}/download"
            return jsonify(entry)
    return jsonify({"error": "Transcription introuvable."}), 404


@app.route("/api/transcriptions/<transcription_id>/download", methods=["GET"])
def download_transcription(transcription_id):
    with _history_lock:
        history = load_history()
    entry = next((item for item in history if item["id"] == transcription_id), None)
    if entry is None:
        return jsonify({"error": "Transcription introuvable."}), 404
    document_path = os.path.join(TRANSCRIPTIONS_DIR, entry["filename"])
    if not os.path.exists(document_path):
        return jsonify({"error": "Document Word manquant sur le serveur."}), 410
    # send_from_directory empêche les attaques par traversée de chemin
    return send_from_directory(
        TRANSCRIPTIONS_DIR,
        entry["filename"],
        as_attachment=True,
        download_name=entry["filename"],
    )


def _delete_entry(entry):
    document_path = os.path.join(TRANSCRIPTIONS_DIR, entry.get("filename", ""))
    if os.path.exists(document_path):
        os.unlink(document_path)


@app.route("/api/transcriptions/<transcription_id>", methods=["DELETE"])
def delete_transcription(transcription_id):
    with _history_lock:
        history = load_history()
        entry = next((item for item in history if item["id"] == transcription_id), None)
        if entry is None:
            return jsonify({"error": "Transcription introuvable."}), 404
        _delete_entry(entry)
        history = [item for item in history if item["id"] != transcription_id]
        save_history(history)
    return "", 204


@app.route("/api/transcriptions", methods=["DELETE"])
def clear_transcriptions():
    with _history_lock:
        history = load_history()
        for entry in history:
            _delete_entry(entry)
        save_history([])
    return "", 204


@app.errorhandler(413)
def file_too_large(_error):
    return (
        jsonify(
            {"error": f"Fichier trop volumineux (limite : {MAX_CONTENT_LENGTH_MB} Mo)."}
        ),
        413,
    )


@app.errorhandler(404)
def not_found(_error):
    return jsonify({"error": "Ressource introuvable."}), 404


@app.errorhandler(500)
def internal_error(_error):
    logger.exception("Erreur serveur")
    return jsonify({"error": "Erreur interne du serveur."}), 500


@app.route("/")
def home():
    return jsonify(
        {
            "service": "Chatbot Vocal",
            "status": "actif",
            "endpoints": [
                "GET /health",
                "POST /api/transcribe",
                "GET /api/transcriptions",
                "GET /api/transcriptions/<id>/download",
            ],
        }
    )


if __name__ == "__main__":
    debug = os.environ.get("FLASK_DEBUG", "0") == "1"
    host = os.environ.get("HOST", "0.0.0.0")
    port = int(os.environ.get("PORT", "5001"))
    logger.info("Démarrage du serveur sur http://%s:%s (debug=%s)", host, port, debug)
    app.run(debug=debug, host=host, port=port)
